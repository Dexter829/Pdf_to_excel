import streamlit as st
import openpyxl
import re
import fitz  
from docx import Document

def ex_info(cv_text):
  #defing my own regular expression for finding the mail id and other revlevant Inforamtion
    email_list = re.findall(r"[a-z0-9\.\-+_]+@[a-z0-9\-]+\.[a-z]+", cv_text, re.IGNORECASE)
    contact_list = re.findall(r"\d{3}-\d{3}-\d{4}|\d{10}", cv_text)
    full_text = cv_text.strip() 
    return email_list, contact_list, full_text


def extract_text_from_pdf(file):
    pdf_bytes = file.read()
    pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = ""
    for page_num in range(len(pdf_document)):
        page = pdf_document.load_page(page_num)
        text += page.get_text()
    return text

def extract_text_from_docx(file):
    docx_document = Document(file)
    text = ""
    for paragraph in docx_document.paragraphs:
        text += paragraph.text
    return text

def download_xlsx(data):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Email", "Contact Number", "Full Text"])  # Header row

    for row in data:
        sheet.append(row)

    wb.save("extracted_info.xlsx")
    st.success("Extracted information saved to extracted_info.xlsx")

st.title("CV Information Extractor Assignment ")
uploaded_file = st.file_uploader("Upload your CVs here!! (PDF, DOCX, or text files)", type=["pdf", "docx", "txt"], accept_multiple_files=True)

if uploaded_file is not None:
    extracted_data = []
    for file in uploaded_file: 
        if file.type == "application/pdf":
            cv_text = extract_text_from_pdf(file)
        elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            cv_text = extract_text_from_docx(file)
        else:
            cv_text = file.read().decode("utf-8")
        
        email_list, contact_list, full_text = ex_info(cv_text)
        extracted_data.append([", ".join(email_list), ", ".join(contact_list), full_text])

    if extracted_data:
        download_button = st.button("Download Extracted Information (.xlsx)")
        if download_button:
            download_xlsx(extracted_data)
